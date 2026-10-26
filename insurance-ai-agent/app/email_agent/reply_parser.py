"""Understand incoming email.

A) INSURER REPLY: Excel/PDF attachment (or body table) that maps PASSPORT NO -> policy no.
   Allianz reply fills E-TRAVEL POLICY NO / ACS POLICY NO; LOLC fills INSURANCE POLICY NO
   (HODTSTxxxxx — which the EXTEND sheet later needs as 'existing policy'). We parse both.
B) EMBASSY APPOINTMENT: body contains appointment/interview wording + a date + a passport
   no (or a known applicant name). We extract {passport_no, appointment_date}.
"""
from __future__ import annotations
import re
from pathlib import Path
from ..dates_engine import parse_any

_POLICY_HINTS = ("POLICY NO", "POLICY NUMBER", "E-TRAVEL", "ACS POLICY", "HODTST")

def _rows_from_xlsx(path: str) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title.startswith("_"): continue
        it = ws.iter_rows(values_only=True)
        header = None
        for r in it:
            cells = [str(c or "").strip() for c in r]
            up = [c.upper() for c in cells]
            if header is None:
                if any("PASSPORT" in c for c in up): header = up
                continue
            rows.append(dict(zip(header, cells)))
    return rows

def _text_from_pdf(path: str) -> str:
    from ..ocr.engine import extract_text
    return extract_text(path)["text"]

def parse_policy_reply(attachments: list[str], body: str = "") -> dict[str, str]:
    """Return {PASSPORT_NO: POLICY_NO} from attachments/body (best effort, empty if none)."""
    mapping: dict[str, str] = {}
    for ap in attachments or []:
        p = Path(ap)
        try:
            if p.suffix.lower() in {".xlsx", ".xlsm"}:
                for row in _rows_from_xlsx(str(p)):
                    passport = next((v for k, v in row.items() if "PASSPORT" in k and v), "")
                    policy = next((v for k, v in row.items()
                                   if any(h in k for h in _POLICY_HINTS) and v), "")
                    if passport and policy:
                        mapping[passport.strip().upper()] = policy.strip()
            elif p.suffix.lower() == ".pdf":
                txt = _text_from_pdf(str(p)).upper()
                for passport, policy in re.findall(r"\b([NP]\d{7,8})\b[^\n]{0,60}?\b(HODTST\d+|[A-Z]{2,}\d{5,})\b", txt):
                    mapping[passport] = policy
                for policy, passport in re.findall(r"\b(HODTST\d+)\b[^\n]{0,120}?\b([NP]\d{7,8})\b", txt):
                    mapping.setdefault(passport, policy)
        except Exception:
            continue
    if body:
        for passport, policy in re.findall(r"\b([NP]\d{7,8})\b[^\n]{0,60}?\b(HODTST\d+)\b", body.upper()):
            mapping.setdefault(passport, policy)
    return mapping

_MONTHS = (r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|"
           r"Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)")

def parse_appointment(subject: str, body: str, keywords: list[str]) -> dict | None:
    blob = f"{subject}\n{body}"
    if not any(k.lower() in blob.lower() for k in keywords):
        return None
    m = re.search(r"\b([NP]\d{7,8})\b", blob.upper())
    date_str = ""
    m2 = re.search(rf"\b\d{{1,2}}\s+{_MONTHS}\s+\d{{4}}\b|\b{_MONTHS}\s+\d{{1,2}},?\s+\d{{4}}\b|"
                   r"\b\d{1,2}[./\-]\d{1,2}[./\-]\d{4}\b", blob, flags=re.I)
    if m2: date_str = m2.group(0)
    d = parse_any(date_str) if date_str else None
    if not d: return None
    return {"passport_no": m.group(1) if m else "",
            "appointment_date": d.isoformat(), "raw_date": date_str}
