"""Excel generation — one file per (company, NEW/EXTEND, day).

- Exact insurer column orders + header spellings (constants.py)
- SALUTATION / GENDER / CITY written as values AND protected by Excel
  dropdown (Data Validation) lists on a hidden `_LISTS` sheet.
- Passport no / phone / agent codes are TEXT (leading zeros survive).
- Dates are TEXT in the insurer's required format.
- Every visible row/column has a full cell outline; columns auto-fit all
  headers and values without vertical text wrapping.
"""
from __future__ import annotations
from datetime import date, datetime
from pathlib import Path
import json
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .constants import COMPANIES, SALUTATIONS, GENDERS


def _sheet_date(value, company: str, column: str) -> str:
    """Render dates using the exact per-column style in the supplied insurer sheets.

    Allianz uses M/D/YYYY for policy dates but DD.MM.YYYY for birthdays.
    LOLC uses DD.MM.YYYY throughout. Do not use a single configurable date
    separator for these mixed insurer formats.
    """
    if not value:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        raw = str(value).strip()
        try:
            value = date.fromisoformat(raw[:10])
        except ValueError:
            patterns = (["%m/%d/%Y", "%m/%d/%y"] if company == "ALLIANZ" and column != "BIRTHDAY" else [])
            patterns += ["%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%d/%m/%y"]
            for pattern in patterns:
                try:
                    value = datetime.strptime(raw, pattern).date()
                    break
                except ValueError:
                    continue
            else:
                return raw
    if company == "ALLIANZ" and column != "BIRTHDAY":
        return f"{value.month}/{value.day}/{value.year}"
    return f"{value.day:02d}.{value.month:02d}.{value.year}"


def _sheet_date_value(value, company: str, column: str):
    """Parse an edited date cell according to that insurer's sample format."""
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        patterns = (["%m/%d/%Y", "%m/%d/%y"] if company == "ALLIANZ" and column != "BIRTHDAY" else [])
        patterns += ["%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%d/%m/%y"]
        for pattern in patterns:
            try:
                return datetime.strptime(raw, pattern).date()
            except ValueError:
                continue
    return None


def _row_for(a: dict, company: str, kind: str, cfg_defaults: dict) -> dict:
    comp = COMPANIES[company]
    d = lambda v, col: _sheet_date(v, company, col)
    base = {
        "ISSUE DATE": d(a.get("issue_date"), "ISSUE DATE"), "SALUTATION": a.get("salutation", ""),
        "GENDER": a.get("gender", ""),
        # insurer convention: FIRST NAME = surname, LAST NAME = given names
        "FIRST NAME": a.get("surname", ""), "LAST NAME": a.get("given_names", ""),
        "PASSPORT NO": a.get("passport_no", ""), "BIRTHDAY": d(a.get("birthday"), "BIRTHDAY"),
        "DEPATURE DATE": d(a.get("departure_date"), "DEPATURE DATE"), "ARRIVAL DATE": d(a.get("arrival_date"), "ARRIVAL DATE"),
        "NO.": a.get("addr_no", ""), "STREET": a.get("addr_street", ""), "CITY": a.get("city", ""),
        "CONTACT NO": a.get("telephone", ""),
        "INSURANCE POLICY NO": a.get("existing_policy_no", ""),
        "E-TRAVEL POLICY NO": a.get("existing_policy_no", "") if kind == "EXTEND" else "",
        "ACS POLICY NO": "", "CUSTOMER PIN": "",
    }
    if company == "ALLIANZ":
        dd = {**comp["defaults"], **cfg_defaults}
        base.update({
            "STATUS": dd.get("STATUS_NEW" if kind == "NEW" else "STATUS_EXTEND", "NEW"),
            "TERRITORIAL AREA": dd.get("TERRITORIAL AREA", ""),
            "PRODUCT": dd.get("PRODUCT", ""), "TOTAL PREMIUM": dd.get("TOTAL PREMIUM", ""),
            "AGENT NAME": dd.get("AGENT NAME", ""), "ACS AGENT CODE": dd.get("ACS AGENT CODE", ""),
            "I CODE": dd.get("I CODE", ""), "INDIVIDUAL CONTRACT": dd.get("INDIVIDUAL CONTRACT", ""),
        })
    if company == "LOLC" and kind == "EXTEND":
        base["ISSUE DATE"] = ""  # blank in LOLC EXTEND sample — insurer fills on issue
    try:
        overrides = json.loads(a.get("sheet_overrides_json") or "{}")
    except (TypeError, ValueError):
        overrides = {}
    if isinstance(overrides, dict):
        allowed = set(comp["columns_new"] if kind == "NEW" else comp["columns_extend"])
        for column, value in overrides.items():
            if column in allowed:
                if column in {"ISSUE DATE", "BIRTHDAY", "DEPATURE DATE", "ARRIVAL DATE"}:
                    base[column] = _sheet_date(value, company, column)
                else:
                    base[column] = value
    if isinstance(base.get("TOTAL PREMIUM"), str):
        try:
            amount = float(base["TOTAL PREMIUM"].replace(",", ""))
            base["TOTAL PREMIUM"] = int(amount) if amount.is_integer() else amount
        except (ValueError, AttributeError):
            pass
    return base


def _format_visible_sheet(ws, cols: list[str]) -> None:
    """Apply full cell borders and horizontally readable auto-fit widths."""
    thin = Side(style="thin", color="9AA8B7")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                min_col=1, max_col=len(cols)), start=1):
        for cell in row:
            cell.border = grid
            font = copy(cell.font)
            font.color = "000000"
            cell.font = font
            cell.alignment = Alignment(horizontal="center" if row_idx == 1 else "left",
                                       vertical="center", wrap_text=False, text_rotation=0)
    # Fit each column to its longest header/value, including defaults repeated
    # in every Allianz row. Text stays horizontal; long fields widen the column.
    for col_idx, _name in enumerate(cols, start=1):
        letter = get_column_letter(col_idx)
        longest = 0
        for cell in ws[letter]:
            if cell.value is not None:
                text = str(cell.value)
                longest = max(longest, *(len(line) for line in text.splitlines()), 0)
        ws.column_dimensions[letter].width = min(255, max(12, longest + 2))
    ws.sheet_view.showGridLines = True


def generate_excel(applicants: list[dict], company: str, kind: str,
                   cities: list[str], out_path: str | Path,
                   cfg_defaults: dict | None = None) -> Path:
    comp = COMPANIES[company]
    cols = comp["columns_new"] if kind == "NEW" else comp["columns_extend"]
    cfg_defaults = cfg_defaults or {}
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, size=9)
    for a in applicants:
        row = _row_for(a, company, kind, cfg_defaults)
        ws.append([row.get(c, "") for c in cols])
    # text-format sensitive columns
    for idx, c in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        if c in {"PASSPORT NO", "CONTACT NO", "ACS AGENT CODE", "I CODE", "NO.",
                 "INSURANCE POLICY NO", "E-TRAVEL POLICY NO", "ACS POLICY NO", "CUSTOMER PIN"}:
            for cell in ws[letter][1:]:
                cell.number_format = "@"
                if cell.value not in ("", None): cell.value = str(cell.value)
    _format_visible_sheet(ws, cols)
    # ---- hidden dropdown lists + validation (SALUTATION / GENDER / CITY) ----
    lists = wb.create_sheet("_LISTS")
    for i, s in enumerate(SALUTATIONS, 1): lists.cell(row=i, column=1, value=s)
    for i, g in enumerate(GENDERS, 1): lists.cell(row=i, column=2, value=g)
    for i, ct in enumerate(cities, 1): lists.cell(row=i, column=3, value=ct)
    lists.sheet_state = "veryHidden"
    last = max(len(applicants) + 1, 2)
    def add_dv(col_name, formula):
        if col_name not in cols: return
        letter = get_column_letter(cols.index(col_name) + 1)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        dv.error = f"Please choose {col_name} from the dropdown list."
        dv.errorTitle = "Invalid entry"; ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{max(last, 500)}")
    add_dv("SALUTATION", f"_LISTS!$A$1:$A${len(SALUTATIONS)}")
    add_dv("GENDER", f"_LISTS!$B$1:$B${len(GENDERS)}")
    add_dv("CITY", f"_LISTS!$C$1:$C${len(cities)}")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    out = Path(out_path); out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def batch_filename(company: str, kind: str, day: date) -> str:
    return f"{company}_{kind}_{day.isoformat()}.xlsx"
