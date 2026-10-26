from datetime import date
from pathlib import Path
import yaml
from flask import Flask
import pytest

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture
def ui_app(tmp_path):
    from app.models import db
    from app.city_matcher import CityMatcher
    from app.routes import bp
    app=Flask('ui-test',template_folder=str(ROOT/'app'/'templates'))
    app.config.update(SECRET_KEY='test-ui',SQLALCHEMY_DATABASE_URI=f"sqlite:///{tmp_path/'ui.sqlite'}",
        SQLALCHEMY_TRACK_MODIFICATIONS=False,MAX_CONTENT_LENGTH=60*1024*1024,
        APP_CFG=yaml.safe_load((ROOT/'config.yaml').read_text()),BASE=ROOT,
        UPLOAD_DIR=tmp_path/'uploads',OUTPUT_DIR=tmp_path/'output',INBOX_DIR=tmp_path/'inbox',
        MAIN_SYSTEM_API_KEY='test-key')
    for k in ('UPLOAD_DIR','OUTPUT_DIR','INBOX_DIR'): Path(app.config[k]).mkdir(parents=True,exist_ok=True)
    db.init_app(app); app.config['CITY_MATCHER']=CityMatcher(ROOT/'data'/'cities.txt'); app.register_blueprint(bp)
    with app.app_context(): db.create_all()
    yield app
    with app.app_context(): db.drop_all()


def test_allianz_defaults_repeat_on_every_new_and_extend_batch_row(ui_app, tmp_path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from app.constants import COMPANIES, ALLIANZ_COLUMNS
    from app.excel_gen import generate_excel
    cfg=yaml.safe_load((ROOT/'config.yaml').read_text())['companies']['ALLIANZ']
    expected={
        'TERRITORIAL AREA':'Worldwide E', 'TOTAL PREMIUM':15500,
        'AGENT NAME':'DDP JAYASURIYA', 'ACS AGENT CODE':'0016938',
        'I CODE':'0017183',
        'INDIVIDUAL CONTRACT':'Overseas Emergency Accident and Sickness Medical Expenses-50 000  Deductible 100',
    }
    applicants=[{'surname':f'PERSON{i}','given_names':'TEST APPLICANT','passport_no':f'N1234567{i}',
        'existing_policy_no':f'LP000000000{i+1}','telephone':'0771234567','salutation':'Mr.','gender':'Male','issue_date':date(2026,10,22),
        'departure_date':date(2026,10,22),'arrival_date':date(2027,4,22),'birthday':date(1990,1,1),
        'addr_no':'10','addr_street':'MAIN ROAD','city':'Dankotuwa'} for i in range(2)]
    for kind, expected_status in [('NEW','NEW'),('EXTEND','EXTEND')]:
        out=tmp_path/f'ALLIANZ_{kind}.xlsx'
        generate_excel(applicants,'ALLIANZ',kind,['Dankotuwa'],out,
            {**cfg.get('defaults',{}),'date_sep':cfg.get('date_sep','/')})
        ws=load_workbook(out,data_only=True).active
        assert ws.max_row==3
        for i,row in enumerate(ws.iter_rows(min_row=2,values_only=True)):
            cells=dict(zip(ALLIANZ_COLUMNS,row))
            assert cells['STATUS']==expected_status
            assert (cells['E-TRAVEL POLICY NO'] or '')==(f'LP000000000{i+1}' if kind=='EXTEND' else '')
            assert {key:cells[key] for key in expected}==expected
        contract_col=ALLIANZ_COLUMNS.index('INDIVIDUAL CONTRACT')+1
        contract_letter=get_column_letter(contract_col)
        assert ws.column_dimensions[contract_letter].width >= len(expected['INDIVIDUAL CONTRACT'])+2
        for row in ws.iter_rows(min_row=1,max_row=ws.max_row,min_col=1,max_col=len(ALLIANZ_COLUMNS)):
            for cell in row:
                assert all(getattr(cell.border,side).style=='thin' for side in ('left','right','top','bottom'))
                assert not cell.alignment.wrap_text and cell.alignment.textRotation==0
                assert cell.font.color and cell.font.color.type=='rgb' and cell.font.color.rgb[-6:]=='000000'


def test_allianz_defaults_are_rendered_in_new_and_extend_sheet_previews(ui_app):
    c=ui_app.test_client()
    expected=['Worldwide E','15500','DDP JAYASURIYA','0016938','0017183',
        'Overseas Emergency Accident and Sickness Medical Expenses-50 000  Deductible 100']
    for kind in ('NEW','EXTEND'):
        html=c.get(f'/insurance/{kind}/SINGLE').get_data(as_text=True)
        assert 'id="format-preview"' in html and 'Selected company’s insurer-sheet preview' in html
        for value in expected: assert value in html,value
        assert 'Exact uploaded column names and order' in html
        assert 'white-space:nowrap' in html and 'border:1px solid #cbd5df' in html
        assert 'Worldwide E' in html and kind in html
        assert 'width:max-content;min-width:100%' in html


def test_insurance_entry_pages_preview_correct_company_formats(ui_app):
    c=ui_app.test_client()
    for path in ('/insurance/NEW/SINGLE','/insurance/NEW/BATCH','/insurance/EXTEND/SINGLE','/insurance/EXTEND/BATCH'):
        r=c.get(path); assert r.status_code==200,path
        text=r.get_data(as_text=True)
        assert 'INSURANCE COMPANY' in text and 'id="format-preview"' in text
        assert 'url.searchParams.set' in text and 'url.searchParams.delete' in text
    new=c.get('/insurance/NEW/SINGLE').get_data(as_text=True)
    assert 'DEPATURE DATE' in new and 'CONTACT NO' in new
    ext=c.get('/insurance/EXTEND/BATCH?company=LOLC').get_data(as_text=True)
    assert 'INSURANCE POLICY NO' in ext and 'batch_files' in ext and 'policy documents' in ext


def test_approved_batch_groups_appear_in_email_queue_and_prepare(ui_app):
    from app.models import db,Applicant,Batch
    app=ui_app
    with app.app_context():
        for i in range(2):
            db.session.add(Applicant(company='LOLC',kind='NEW',entry_mode='BATCH',status='APPROVED',
                surname=f'TEST{i}',given_names='CANDIDATE',passport_no=f'N1234567{i}',
                salutation='Mr.',gender='Male',birthday=date(1990,1,1),issue_date=date(2026,10,22),
                departure_date=date(2026,10,22),arrival_date=date(2027,4,22),city='Dankotuwa'))
        db.session.commit()
    c=app.test_client(); page=c.get('/emails'); assert page.status_code==200
    html=page.get_data(as_text=True); assert 'Approved · waiting for email preparation' in html
    r=c.post('/batches',data={'company':'LOLC','kind':'NEW','submission_mode':'BATCH',
        'batch_date':'2026-10-22','ids':['1','2']})
    assert r.status_code==302
    with app.app_context():
        b=Batch.query.first(); assert b.submission_mode=='BATCH' and Path(b.excel_path).exists()
        assert len(b.applicants)==2
    assert c.get('/emails').status_code==200


def test_email_queue_filters_and_lists_single_applicants_by_company_type_and_mode(ui_app):
    from app.models import db, Applicant
    app=ui_app
    with app.app_context():
        records=[
            ('ALLIANZ','NEW','SINGLE','ALPHA','N11111111'),
            ('LOLC','NEW','SINGLE','BETA','N22222222'),
            ('ALLIANZ','EXTEND','SINGLE','GAMMA','N33333333'),
            ('LOLC','EXTEND','BATCH','DELTA','N44444444'),
        ]
        for company,kind,mode,name,passport in records:
            db.session.add(Applicant(company=company,kind=kind,entry_mode=mode,status='APPROVED',
                surname=name,given_names='QUEUE TEST',passport_no=passport,salutation='Mr.',gender='Male',
                birthday=date(1990,1,1),telephone='0701234567',city='Colombo',
                issue_date=date(2026,10,22),departure_date=date(2026,10,22),arrival_date=date(2027,4,22),
                embassy_appointment_date=date(2026,10,8) if kind=='EXTEND' else None))
        db.session.commit()
    client=app.test_client()
    page=client.get('/emails').get_data(as_text=True)
    for value in ('ALLIANZ','LOLC','NEW','EXTEND','SINGLE','BATCH','ALPHA QUEUE TEST',
                  'BETA QUEUE TEST','GAMMA QUEUE TEST','DELTA QUEUE TEST',
                  'N11111111','N22222222','N33333333','N44444444'):
        assert value in page, value
    filtered=client.get('/emails?company=ALLIANZ&kind=EXTEND&mode=SINGLE').get_data(as_text=True)
    assert 'GAMMA QUEUE TEST' in filtered and 'N33333333' in filtered
    assert 'ALPHA QUEUE TEST' not in filtered and 'BETA QUEUE TEST' not in filtered
    assert 'DELTA QUEUE TEST' not in filtered
    assert 'Prepared email jobs' in filtered  # history remains visible regardless of filters


def test_yahoo_smtp_sender_uses_configured_host_and_attachment_without_network(ui_app,tmp_path,monkeypatch):
    from unittest.mock import patch
    from app.email_agent.smtp_imap import send_email
    attached=tmp_path/'applicant.xlsx';attached.write_bytes(b'test workbook')
    calls={}
    class FakeSMTP:
        def __init__(self,host,port,timeout): calls.update(host=host,port=port,timeout=timeout)
        def __enter__(self): return self
        def __exit__(self,*args): return False
        def starttls(self,context=None): calls['starttls']=bool(context)
        def login(self,address,password): calls['login']=(address,password)
        def send_message(self,message): calls['message']=message
    monkeypatch.setenv('MAIL_ADDRESS','sender@yahoo.com')
    monkeypatch.setenv('MAIL_PASSWORD','test-only-app-password')
    with patch('app.email_agent.smtp_imap.smtplib.SMTP',FakeSMTP):
        send_email(ui_app.config['APP_CFG']['email'],'insurer@example.com','Test subject','Test body',[attached])
    assert calls['host']=='smtp.mail.yahoo.com' and calls['port']==587
    assert calls['starttls'] and calls['login']==('sender@yahoo.com','test-only-app-password')
    message=calls['message']
    assert message['To']=='insurer@example.com' and message['Subject']=='Test subject'
    parts=list(message.iter_attachments())
    assert len(parts)==1 and parts[0].get_filename()=='applicant.xlsx'
    assert parts[0].get_payload(decode=True)==b'test workbook'


def test_email_auto_send_setting_sends_after_prepare_only_when_enabled(ui_app):
    from app.models import db,Applicant,Batch
    from unittest.mock import patch
    app=ui_app
    with app.app_context():
        a=Applicant(company='ALLIANZ',kind='NEW',entry_mode='SINGLE',status='APPROVED',
            surname='AUTO',given_names='TEST',passport_no='N76543210',salutation='Mr.',gender='Male',
            birthday=date(1990,1,1),telephone='0771234567',issue_date=date(2026,10,22),
            departure_date=date(2026,10,22),arrival_date=date(2027,4,22),city='Dankotuwa')
        db.session.add(a); db.session.commit(); aid=a.id
    c=app.test_client(); c.post('/emails/auto-send',data={'enabled':'1'}); captured=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''): captured.append((to,body,len(attachments),cc))
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        response=c.post('/batches',data={'company':'ALLIANZ','kind':'NEW','submission_mode':'SINGLE',
            'batch_date':'2026-10-22','ids':[str(aid)]})
    assert response.status_code==302 and len(captured)==1
    assert captured[0][0]=='chamikafernando456@gmail.com' and captured[0][2]==1
    assert captured[0][3]=='chamikaf@allianz.lk'
    assert 'warunic@allianz.lk' not in captured[0][3]
    with app.app_context():
        b=Batch.query.first(); assert b.status=='SENT'
        assert db.session.get(Applicant,aid).status=='AWAITING_REPLY'


def _approval_form(company='LOLC',kind='NEW',surname='TEST',passport='N76543210'):
    import json
    row={'ISSUE DATE':'2026-10-22','SALUTATION':'Mr.','GENDER':'Male','FIRST NAME':surname,
        'LAST NAME':'CANDIDATE','PASSPORT NO':passport,'BIRTHDAY':'1990-01-01',
        'DEPATURE DATE':'2026-10-22','ARRIVAL DATE':'2027-04-22','NO.':'10/A',
        'STREET':'TEST ROAD','CITY':'Dankotuwa'}
    if company=='ALLIANZ': row.update({'STATUS':'NEW' if kind=='NEW' else 'EXTEND',
        'TERRITORIAL AREA':'Worldwide E','PRODUCT':'Single Trip Gold Plan - USD 50,000 Coverage (per person)',
        'TOTAL PREMIUM':'15500','AGENT NAME':'DDP JAYASURIYA','ACS AGENT CODE':'0016938',
        'I CODE':'0017183','INDIVIDUAL CONTRACT':'Overseas Emergency Accident and Sickness Medical Expenses-50 000  Deductible 100',
        'CONTACT NO':'0771234567'})
    elif kind=='EXTEND': row['INSURANCE POLICY NO']='LP123456'
    return {'sheet_row_json':json.dumps(row),'final_check':'1','approve':'1'}


def test_extend_email_body_displays_departure_date_not_arrival_date(ui_app):
    from types import SimpleNamespace
    from app.routes import _batch_email_body
    applicant=SimpleNamespace(surname='TEST',given_names='TRAVELLER',
                              departure_date=date(2026,10,15),arrival_date=date(2027,4,15))
    for company in ('ALLIANZ','LOLC'):
        batch=SimpleNamespace(company=company,kind='EXTEND')
        body=_batch_email_body(batch,[applicant])
        assert 'OCTOBER 15, 2026' in body, body
        assert 'APRIL 15, 2027' not in body, body
        assert 'departure date' in body.lower()
        assert 'arrival date' not in body.lower()


def test_auto_send_single_happens_after_approval(ui_app):
    from app.models import db,Applicant,Batch
    from unittest.mock import patch
    app=ui_app
    with app.app_context():
        a=Applicant(company='LOLC',kind='NEW',entry_mode='SINGLE',status='NEEDS_REVIEW',
            surname='TEST',given_names='SINGLE',passport_no='N76543210',salutation='Mr.',gender='Male',
            birthday=date(1990,1,1),issue_date=date(2026,10,22),departure_date=date(2026,10,22),
            arrival_date=date(2027,4,22),city='Dankotuwa')
        db.session.add(a);db.session.commit();aid=a.id
    c=app.test_client();c.post('/emails/auto-send',data={'enabled':'1'});sent=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''):sent.append((to,len(attachments)))
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        c.post(f'/review/{aid}',data=_approval_form(surname='TEST',passport='N76543210'))
    assert sent==[('suthsil73@gmail.com',1)]
    with app.app_context():
        assert Batch.query.first().status=='SENT'
        assert db.session.get(Applicant,aid).status=='AWAITING_REPLY'


def test_review_selected_passport_copy_is_attached_when_auto_sending(ui_app):
    from app.models import db,Applicant
    from unittest.mock import patch
    app=ui_app
    passport=Path(app.config['UPLOAD_DIR'])/'passport.pdf'; passport.write_bytes(b'fake passport copy')
    with app.app_context():
        a=Applicant(company='LOLC',kind='NEW',entry_mode='SINGLE',status='NEEDS_REVIEW',
            surname='COPY',given_names='TEST',passport_no='N76543210',salutation='Mr.',gender='Male',
            birthday=date(1990,1,1),issue_date=date(2026,10,22),departure_date=date(2026,10,22),
            arrival_date=date(2027,4,22),city='Dankotuwa',passport_file_path=str(passport))
        db.session.add(a);db.session.commit();aid=a.id
    c=app.test_client();c.post('/emails/auto-send',data={'enabled':'1'});captured=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''):captured.append(list(attachments))
    form=_approval_form(surname='COPY',passport='N76543210')
    form['attach_passport_copy']=['0','1']
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        c.post(f'/review/{aid}',data=form)
    assert len(captured)==1 and captured[0][1]==str(passport)
    with app.app_context():
        assert db.session.get(Applicant,aid).attach_passport_copy is True


def test_auto_batch_waits_for_all_same_date_approvals(ui_app):
    from app.models import db,Applicant,Batch
    from unittest.mock import patch
    from openpyxl import load_workbook
    app=ui_app
    with app.app_context():
        ids=[]
        for i in range(2):
            a=Applicant(company='LOLC',kind='NEW',entry_mode='BATCH',status='NEEDS_REVIEW',
                surname=f'TEST{i}',given_names='BATCH',passport_no=f'N7654321{i}',salutation='Mr.',gender='Male',
                birthday=date(1990,1,1),issue_date=date(2026,10,22),departure_date=date(2026,10,22),
                arrival_date=date(2027,4,22),city='Dankotuwa')
            db.session.add(a);db.session.flush();ids.append(a.id)
        db.session.commit()
    c=app.test_client();c.post('/emails/auto-send',data={'enabled':'1'});sent=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''):
        sent.append((to,load_workbook(attachments[0],data_only=True).active.max_row))
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        c.post(f'/review/{ids[0]}',data=_approval_form(surname='TEST0',passport='N76543210'))
        assert sent==[]  # wait for the other same-date applicant
        c.post(f'/review/{ids[1]}',data=_approval_form(surname='TEST1',passport='N76543211'))
    assert sent==[('suthsil73@gmail.com',3)]  # header + both applicants
    with app.app_context():
        assert Batch.query.count()==1 and Batch.query.first().submission_mode=='BATCH'


def test_auto_batch_attaches_only_the_passport_selected_for_that_applicant(ui_app):
    from app.models import db,Applicant,Batch
    from unittest.mock import patch
    app=ui_app
    first_passport=Path(app.config['UPLOAD_DIR'])/'first-passport.pdf'
    second_passport=Path(app.config['UPLOAD_DIR'])/'second-passport.pdf'
    first_passport.write_bytes(b'first applicant passport')
    second_passport.write_bytes(b'second applicant passport')
    with app.app_context():
        ids=[]
        for i,(name,path) in enumerate((('FIRST',first_passport),('SECOND',second_passport))):
            applicant=Applicant(company='LOLC',kind='NEW',entry_mode='BATCH',status='NEEDS_REVIEW',
                surname=name,given_names='BATCH TEST',passport_no=f'N7654321{i}',salutation='Mr.',gender='Male',
                birthday=date(1990,1,1),issue_date=date(2026,10,22),departure_date=date(2026,10,22),
                arrival_date=date(2027,4,22),city='Dankotuwa',passport_file_path=str(path))
            db.session.add(applicant);db.session.flush();ids.append(applicant.id)
        db.session.commit()
    client=app.test_client();client.post('/emails/auto-send',data={'enabled':'1'});captured=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''):
        captured.append((list(attachments),[Path(p).read_bytes() for p in attachments]))
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        first_form=_approval_form(surname='FIRST',passport='N76543210')
        first_form['attach_passport_copy']=['0','1']
        client.post(f'/review/{ids[0]}',data=first_form)
        assert captured==[]
        client.post(f'/review/{ids[1]}',data=_approval_form(surname='SECOND',passport='N76543211'))
    assert len(captured)==1
    attachments,contents=captured[0]
    assert len(attachments)==2  # Excel + selected first applicant passport
    assert str(first_passport) in attachments
    assert str(second_passport) not in attachments
    assert b'first applicant passport' in contents and b'second applicant passport' not in contents
    with app.app_context():
        assert db.session.get(Applicant,ids[0]).attach_passport_copy is True
        assert db.session.get(Applicant,ids[1]).attach_passport_copy is False
        assert Batch.query.first().status=='SENT'


def test_review_sheet_overrides_are_used_by_generated_excel(tmp_path):
    import json
    from openpyxl import load_workbook
    from app.excel_gen import generate_excel
    applicant={'surname':'TEST','given_names':'APPLICANT','passport_no':'N12345678',
        'salutation':'Mr.','gender':'Male','birthday':date(1990,1,1),
        'issue_date':date(2026,10,22),'departure_date':date(2026,10,22),
        'arrival_date':date(2027,4,22),'telephone':'0771234567','addr_no':'1',
        'addr_street':'ROAD','city':'Colombo',
        'sheet_overrides_json':json.dumps({'TOTAL PREMIUM':'17250','PRODUCT':'MANUALLY EDITED PRODUCT','STATUS':'CUSTOM'})}
    out=tmp_path/'edited.xlsx'
    generate_excel([applicant],'ALLIANZ','NEW',['Colombo'],out)
    ws=load_workbook(out,data_only=True).active
    values={ws.cell(1,c).value:ws.cell(2,c).value for c in range(1,ws.max_column+1)}
    assert values['TOTAL PREMIUM']==17250
    assert values['PRODUCT']=='MANUALLY EDITED PRODUCT'
    assert values['STATUS']=='CUSTOM'


def test_clear_notifications_button_only_deletes_notification_rows(ui_app):
    from app.models import db, EmailLog, Notification
    app=ui_app
    with app.app_context():
        db.session.add_all([
            Notification(title='Test alert one',body='First',level='info'),
            Notification(title='Test alert two',body='Second',level='warning'),
            EmailLog(direction='IN',frm='reply@example.com',to='agency@example.com',
                     subject='Retained email log',kind='OTHER'),
        ])
        db.session.commit()

    client=app.test_client()
    page=client.get('/notifications')
    assert page.status_code==200
    assert b'Clear notifications' in page.data
    assert b'Applicant records and email logs will be kept.' in page.data

    response=client.post('/notifications/clear')
    assert response.status_code==302
    with app.app_context():
        assert Notification.query.count()==0
        assert EmailLog.query.count()==1
        assert EmailLog.query.first().subject=='Retained email log'


def test_clear_batch_source_documents_preserves_linked_passport_copy(ui_app,tmp_path):
    import json
    from app.models import db, Applicant, ExtractionJob
    app=ui_app
    with app.app_context():
        upload_root=Path(app.config['UPLOAD_DIR'])
        source=upload_root/'bulk-test-clear'
        passport=source/'combined_files'/'001'/'scan-passport.pdf'
        police=source/'combined_files'/'002'/'scan-police.pdf'
        passport.parent.mkdir(parents=True);police.parent.mkdir(parents=True)
        passport.write_bytes(b'test passport copy');police.write_bytes(b'test police source')
        applicant=Applicant(company='ALLIANZ',kind='NEW',entry_mode='BATCH',status='NEEDS_REVIEW',
                            surname='TEST',given_names='APPLICANT',passport_no='N12345678',
                            passport_file_path=str(passport))
        db.session.add(applicant);db.session.flush()
        summary={'passport_files':[{'name':passport.name,'path':str(passport),'confidence':0.99}],
                 'police_files':[{'name':police.name,'path':str(police),'confidence':0.8}]}
        job=ExtractionJob(job_id='testclear123',company='ALLIANZ',kind='NEW',entry_mode='BATCH',
            status='COMPLETE',applicant_ids_json=json.dumps([applicant.id]),
            file_summary_json=json.dumps(summary),source_root=str(source),
            redirect_url='/insurance/NEW/BATCH?company=ALLIANZ&job_id=testclear123')
        db.session.add(job);db.session.commit();aid=applicant.id
    client=app.test_client()
    response=client.post('/extraction/testclear123/clear-files')
    assert response.status_code==302
    with app.app_context():
        applicant=db.session.get(Applicant,aid);job=db.session.get(ExtractionJob,'testclear123')
        retained=Path(applicant.passport_file_path)
        assert job.source_files_cleared is True
        assert not source.exists()
        assert retained.is_file() and retained.read_bytes()==b'test passport copy'
        assert retained.parent.name=='passport-copies'
    page=client.get(response.headers['Location'])
    assert page.status_code==200
    assert b'scan-passport.pdf' in page.data and b'scan-police.pdf' in page.data
    assert b'Source scans were cleared from this batch folder' in page.data


def test_batch_upload_accepts_combined_multi_file_selection(ui_app):
    from io import BytesIO
    from unittest.mock import patch
    from app.models import db, ExtractionJob
    app=ui_app;captured={}
    def fake_dispatch(flask_app,job_id,worker,file_paths,*args):
        captured['job_id']=job_id;captured['file_paths']=file_paths
    data={'company':'ALLIANZ','kind':'NEW','apply_date':'2026-10-24',
          'batch_files':[(BytesIO(b'passport sample'),'passport-scan.pdf'),
                         (BytesIO(b'police sample'),'police-scan.pdf'),
                         (BytesIO(b'application sample'),'phone-form.pdf')]}
    with patch('app.routes._dispatch_extraction',side_effect=fake_dispatch):
        response=app.test_client().post('/upload/bulk',data=data,content_type='multipart/form-data')
    assert response.status_code==302
    assert len(captured['file_paths']['combined_files'])==3
    assert [Path(p).name for p in captured['file_paths']['combined_files']]==[
        'passport-scan.pdf','police-scan.pdf','phone-form.pdf']
    with app.app_context():
        job=db.session.get(ExtractionJob,captured['job_id'])
        assert job.entry_mode=='BATCH' and job.source_root
        assert Path(job.source_root).is_dir()


@pytest.mark.parametrize('company,kind', [('LOLC','NEW'),('LOLC','EXTEND'),('ALLIANZ','NEW'),('ALLIANZ','EXTEND')])
def test_insurer_preview_shows_passport_fields_and_shared_dropdowns(ui_app,company,kind):
    import json
    from app.models import db, Applicant, InsurancePreview
    app=ui_app
    with app.app_context():
        applicant=Applicant(company=company,kind=kind,entry_mode='BATCH',status='NEEDS_REVIEW',
            surname='FERNANDO',given_names='CHAMIKA',passport_no='N12345678',
            birthday=date(1990,2,3),gender='Female',salutation='Ms.',
            issue_date=date(2026,10,24),departure_date=date(2026,11,7),
            arrival_date=date(2027,5,7),city='Colombo')
        db.session.add(applicant);db.session.flush();aid=applicant.id
        key=f'{company}:{kind}:BATCH'
        db.session.add(InsurancePreview(scope_key=key,applicant_ids_json=json.dumps([aid])))
        db.session.commit()
    client=app.test_client()
    page=client.get(f'/insurance/{kind}/BATCH?company={company}')
    assert page.status_code==200
    text=page.get_data(as_text=True)
    for column in ('SALUTATION','GENDER','FIRST NAME','LAST NAME','BIRTHDAY'):
        assert column in text
    assert 'FERNANDO' in text and 'CHAMIKA' in text and '03.02.1990' in text
    assert f'name="salutation_{aid}"' in text and f'name="gender_{aid}"' in text
    assert 'option value="Miss"' in text and 'option value="Female"' in text
    response=client.post('/insurance/save-preview-choices',data={
        'company':company,'kind':kind,'mode':'BATCH','applicant_ids':[str(aid)],
        f'salutation_{aid}':'Miss',f'gender_{aid}':'Female'})
    assert response.status_code==302
    with app.app_context():
        saved=db.session.get(Applicant,aid)
        assert saved.salutation=='Miss' and saved.gender=='Female'


def test_final_review_uses_same_salutation_and_gender_dropdowns(ui_app):
    from app.models import db, Applicant
    app=ui_app
    with app.app_context():
        applicant=Applicant(company='LOLC',kind='NEW',entry_mode='SINGLE',status='NEEDS_REVIEW',
            surname='TEST',given_names='TRAVELLER',passport_no='N76543210',
            birthday=date(1990,2,3),gender='Male',salutation='Mr.',city='Colombo')
        db.session.add(applicant);db.session.commit();aid=applicant.id
    page=app.test_client().get(f'/review/{aid}')
    assert page.status_code==200
    text=page.get_data(as_text=True)
    assert 'data-sheet-column="SALUTATION"' in text
    assert 'data-sheet-column="GENDER"' in text
    assert 'option value="Miss"' in text and 'option value="Female"' in text
