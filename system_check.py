"""Safe end-to-end check using temporary SQLite and mocked outbound email.
Uses the provided sample police reports; never sends a real email or edits live DB."""
from __future__ import annotations
import io, sys, tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch
import fitz
from flask import Flask
from openpyxl import Workbook, load_workbook
import yaml

ROOT = Path('/home/user/insurance-ai-agent')
UPLOADS = Path('/home/user/uploads')
POLICE1 = UPLOADS / 'NUWAN PRIYANKARA POLICE REPORT - Copy.pdf'
POLICE2 = UPLOADS / 'IRESH LOCHANA POLICE REPORT - Copy.pdf'
sys.path.insert(0, str(ROOT))
from app.models import db, Applicant, Batch, EmailLog, Notification
from app.city_matcher import CityMatcher
from app.ocr.mrz import check_digit


def pdf_text(path: Path, lines: list[str]):
    doc = fitz.open(); page = doc.new_page()
    y = 70
    for line in lines:
        page.insert_text((50, y), line, fontsize=11); y += 22
    doc.save(path); doc.close()


def passport_pdf(path: Path, passport_no: str, surname: str, given: str,
                 dob='890118', sex='M'):
    expiry='340117'
    l1=('P<LKA'+surname.upper().replace(' ','<')+'<<'+given.upper().replace(' ','<'))
    l1=(l1+'<'*44)[:44]
    l2=(passport_no+str(check_digit(passport_no))+'LKA'+dob+str(check_digit(dob))+
        sex+expiry+str(check_digit(expiry))+'<'*16)
    l2=(l2+'<'*44)[:44]
    pdf_text(path, ['PASSPORT',f'Surname: {surname}',f'Given Names: {given}',
                    f'Passport No: {passport_no}','Date of Birth: 18/01/1989',
                    f'Sex: {sex}',l1,l2])


def make_app(work: Path):
    app = Flask('isolated-system-check', template_folder=str(ROOT/'app'/'templates'))
    app.config.update(SECRET_KEY='test-only',
        SQLALCHEMY_DATABASE_URI=f"sqlite:///{work/'test.sqlite'}",
        SQLALCHEMY_TRACK_MODIFICATIONS=False, MAX_CONTENT_LENGTH=60*1024*1024,
        APP_CFG=yaml.safe_load((ROOT/'config.yaml').read_text()), BASE=ROOT,
        UPLOAD_DIR=work/'uploads', OUTPUT_DIR=work/'output', INBOX_DIR=work/'inbox')
    for key in ('UPLOAD_DIR','OUTPUT_DIR','INBOX_DIR'): Path(app.config[key]).mkdir(parents=True,exist_ok=True)
    db.init_app(app)
    app.config['CITY_MATCHER'] = CityMatcher(ROOT/'data'/'cities.txt')
    from app.routes import bp
    app.register_blueprint(bp)
    with app.app_context(): db.create_all()
    return app


def filetuple(path: Path): return (io.BytesIO(path.read_bytes()), path.name)


def main():
  with tempfile.TemporaryDirectory(prefix='insurance-system-check-') as td:
    work=Path(td); p1=work/'passport1.pdf'; p2=work/'passport2.pdf'; application=work/'application.pdf'
    passport_pdf(p1,'N10199973','WIJELATH','PEDIGE NUWAN PIYANKARA')
    passport_pdf(p2,'N12345678','RAJAPAKSHA','PATHIRANAGE IRESH LOSHANA')
    pdf_text(application,['Applicant application form','Contact mobile phone: 0771234567',
                          'Travel insurance telephone contact number supplied by applicant.'])
    app=make_app(work); c=app.test_client(); reports=[]
    for url in ('/','/upload','/upload/bulk','/applicants','/batches','/appointments','/notifications','/settings'):
        r=c.get(url); assert r.status_code==200,(url,r.status_code)
    print('PASS web dashboard/pages and settings')
    assert 'Marawila' in c.get('/api/cities?q=Maraw').get_data(as_text=True)

    # Single NEW route: passport MRZ + police report OCR + application phone.
    r=c.post('/upload',data={'company':'LOLC','kind':'NEW','apply_date':'2026-10-22',
        'passport':filetuple(p1),'police':filetuple(POLICE1),'application':filetuple(application)},
        content_type='multipart/form-data')
    assert r.status_code==302,r.status_code
    with app.app_context():
        first=Applicant.query.order_by(Applicant.id.desc()).first(); first_id=first.id
        assert first.passport_no=='N10199973',first.passport_no
        assert first.birthday==date(1989,1,18),first.birthday
        assert first.city=='Dankotuwa',first.city
        assert first.addr_no=='41/A' and 'SAMAGI' in first.addr_street,(first.addr_no,first.addr_street)
        assert first.telephone=='771234567',first.telephone
        assert first.departure_date==date(2026,10,22) and first.arrival_date==date(2027,4,22)
        assert first.salutation=='Mr.' and first.gender=='Male'
    assert c.get(f'/review/{first_id}').status_code==200
    r=c.post(f'/review/{first_id}',data={'company':'LOLC','kind':'NEW','salutation':'Mr.',
        'gender':'Male','surname':'WIJELATH','given_names':'PEDIGE NUWAN PIYANKARA',
        'passport_no':'N10199973','birthday':'1989-01-18','telephone':'771234567',
        'addr_no':'41/A','addr_street':'SAMAGI MAWATHA KATUKENDA WEST','city':'Dankotuwa',
        'issue_date':'2026-10-22','departure_date':'2026-10-22','arrival_date':'2027-04-22',
        'recalc_dates':'1','approve':'1'},follow_redirects=False)
    with app.app_context(): assert db.session.get(Applicant,first_id).status=='APPROVED'
    print('PASS single NEW upload → passport MRZ/date/phone/address extraction → review approval')

    # Bulk police/passport pairing, re-upload, and typed-address path.
    r=c.post('/upload/bulk',data={'company':'ALLIANZ','kind':'NEW','apply_date':'2026-10-22',
        'police_files':[filetuple(POLICE1),filetuple(POLICE2)],
        'passport_files':[filetuple(p1),filetuple(p2)]},content_type='multipart/form-data')
    assert r.status_code==302,r.status_code
    with app.app_context():
        bulk=Applicant.query.order_by(Applicant.id.desc()).limit(2).all()
        bycity={a.city:a for a in bulk}; assert set(bycity)=={'Dankotuwa','Marawila'},[(a.city,a.addr_no) for a in bulk]
        bulk_ids=[a.id for a in bulk]
        row=bycity['Marawila']; row.status='EXCEL_GENERATED'; db.session.commit(); rep_id=row.id
    r=c.post(f'/review/{rep_id}/police',data={'police':filetuple(POLICE1)},content_type='multipart/form-data')
    assert r.status_code==302
    with app.app_context():
        row=db.session.get(Applicant,rep_id); assert (row.addr_no,row.city,row.status)==('41/A','Dankotuwa','NEEDS_REVIEW'),(row.addr_no,row.city,row.status)
    second_id=next(i for i in bulk_ids if i!=rep_id)
    r=c.post(f'/review/{second_id}/readdress',data={'address_raw':'NO 203/C, ANASLI HERATH MAWATHA, MUDUKATUWA, MARAWILA, SRI LANKA'})
    with app.app_context():
        row=db.session.get(Applicant,second_id); assert (row.addr_no,row.city)==('203/C','Marawila')
    print('PASS bulk uploads + automatic city match + approved re-upload reset + pasted-address split')

    # Four insurer/workflow exports, each with two same-day/same-appointment applicants.
    issue=date(2026,10,22); appointment=date(2026,10,16)
    groups=[('ALLIANZ','NEW'),('LOLC','NEW'),('ALLIANZ','EXTEND'),('LOLC','EXTEND')]
    with app.app_context():
        for company,kind in groups:
            for j in range(2):
                pp=f'N{(20000000+len(Applicant.query.all())*10+j):08d}'
                a=Applicant(company=company,kind=kind,status='APPROVED',salutation='Mr.',gender='Male',
                    surname=f'TEST{j+1}',given_names='SAMPLE APPLICANT',passport_no=pp,
                    birthday=date(1990,1,18),telephone='771234567',addr_no='10/A',addr_street='TEST STREET',city='Dankotuwa',
                    issue_date=issue,existing_policy_no='HODTST47813' if kind=='EXTEND' else '',
                    embassy_appointment_date=appointment if kind=='EXTEND' else None,
                    departure_date=issue if kind=='NEW' else date(2026,10,30),
                    arrival_date=date(2027,4,22) if kind=='NEW' else date(2027,4,30))
                db.session.add(a)
        db.session.commit()
    batches={}
    for company,kind in groups:
        groupdate=issue if kind=='NEW' else appointment
        r=c.post('/batches',data={'company':company,'kind':kind,'batch_date':groupdate.isoformat()})
        assert r.status_code==302
        with app.app_context():
            b=Batch.query.filter_by(company=company,kind=kind,batch_date=groupdate).order_by(Batch.id.desc()).first()
            assert b and Path(b.excel_path).exists(),(company,kind)
            batches[(company,kind)]=b.id; path=Path(b.excel_path)
            wb=load_workbook(path,data_only=True); ws=wb.active
            header=[v.value for v in ws[1]]; rows=list(ws.iter_rows(min_row=2,values_only=True))
            assert len(rows)>=2,(company,kind,len(rows))
            assert ws.max_row==len(rows)+1
            idx={name:i for i,name in enumerate(header)}
            exp_dep='10/22/2026' if company=='ALLIANZ' and kind=='NEW' else ('22/10/2026' if kind=='NEW' else ('10/30/2026' if company=='ALLIANZ' else '30/10/2026'))
            exp_arr='04/22/2027' if company=='ALLIANZ' and kind=='NEW' else ('22/04/2027' if kind=='NEW' else ('04/30/2027' if company=='ALLIANZ' else '30/04/2027'))
            assert rows[0][idx['DEPATURE DATE']]==exp_dep,(company,kind,rows[0][idx['DEPATURE DATE']],exp_dep)
            assert rows[0][idx['ARRIVAL DATE']]==exp_arr,(company,kind,rows[0][idx['ARRIVAL DATE']],exp_arr)
            assert len(ws.data_validations.dataValidation)>=3
            if company=='LOLC' and kind=='EXTEND': assert rows[0][idx['ISSUE DATE']] in ('',None)
    print('PASS all four Excel formats, grouping, dates, and dropdown validations')

    # SMTP is mocked so no real message goes out. Verify error handling, recipients, copy options and bodies.
    sent=[]
    def fake_send(cfg,to,subject,body,attachments,cc=''):
        assert all(Path(x).exists() for x in attachments)
        sent.append((to,subject,body,list(attachments)))
    first_batch=batches[('LOLC','NEW')]
    with patch('app.email_agent.smtp_imap.send_email',side_effect=RuntimeError('test-only SMTP failure')):
        r=c.post(f'/send/{first_batch}',data={})
        assert r.status_code==302
    with app.app_context(): assert db.session.get(Batch,first_batch).status=='DRAFT'
    # Only the mocked sender below is used; no outbound email is permitted in this check.
    with patch('app.email_agent.smtp_imap.send_email',side_effect=fake_send):
        for key,bid in batches.items():
            if key==('LOLC','NEW'):
                r=c.post(f'/send/{bid}',data={'passport_attachments':(io.BytesIO(b'test passport file'),'test-passport.pdf')},content_type='multipart/form-data')
            else: r=c.post(f'/send/{bid}',data={})
            assert r.status_code==302
    assert len(sent)==4,len(sent)
    bykey={}
    for item in sent:
        recipient,subject,body,attachments=item
        company='ALLIANZ' if recipient=='chamikafe@allianz.lk' else 'LOLC'
        kind='NEW' if body.startswith('Good Day') else 'EXTEND'
        bykey[(company,kind)]=item
    assert bykey[('ALLIANZ','NEW')][0]=='chamikafe@allianz.lk'
    assert 'Insurance Application' in bykey[('ALLIANZ','NEW')][2]
    assert bykey[('LOLC','NEW')][0]=='suthsil73@gmail.com' and 'passport copies' in bykey[('LOLC','NEW')][2].lower() and len(bykey[('LOLC','NEW')][3])==2
    assert 'Dear Allianz Insurance Team' in bykey[('ALLIANZ','EXTEND')][2] and 'OCTOBER 30, 2026' in bykey[('ALLIANZ','EXTEND')][2]
    assert 'Dear Team' in bykey[('LOLC','EXTEND')][2] and 'OCTOBER 30, 2026' in bykey[('LOLC','EXTEND')][2]
    with app.app_context():
        assert all(db.session.get(Batch,bid).status=='SENT' for bid in batches.values())
        assert EmailLog.query.filter_by(direction='OUT').count()==4
    print('PASS all email compositions/recipients/optional copy attach + send-failure guard (SMTP mocked; zero real emails)')

    # Simulated inbound insurer reply and embassy appointment (no Yahoo IMAP login).
    reply=work/'inbox'/'policy_reply.xlsx'; reply.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); ws=wb.active; ws.append(['PASSPORT NO','INSURANCE POLICY NO']); ws.append(['N20000000','HODTST99001']); wb.save(reply)
    with app.app_context():
        src=Applicant(company='LOLC',kind='NEW',status='AWAITING_REPLY',surname='INBOX',given_names='TEST',passport_no='N20000000',issue_date=issue,departure_date=issue,arrival_date=date(2027,4,22))
        apptsrc=Applicant(company='LOLC',kind='NEW',status='POLICY_RECEIVED',surname='APPT',given_names='TEST',passport_no='N76543210',new_policy_no='HODTST99010',issue_date=issue,departure_date=issue,arrival_date=date(2027,4,22))
        db.session.add_all([src,apptsrc]); db.session.commit(); inbox_id=src.id; appt_id=apptsrc.id
    from app.inbox_service import process_inbox
    msg={'from':'test-insurer@example.test','to':'transarabian@yahoo.com','subject':'Policy issued','body':'','message_id':'test-policy','attachments':[str(reply)]}
    with patch('app.email_agent.smtp_imap.fetch_unseen',return_value=[msg]):
        with app.app_context(): result=process_inbox()
    with app.app_context():
        row=db.session.get(Applicant,inbox_id); assert row.new_policy_no=='HODTST99001' and row.status=='POLICY_RECEIVED'
    assert result['policies']>=1
    apptmsg={'from':'test-embassy@example.test','to':'transarabian@yahoo.com','subject':'Embassy interview appointment','body':'Passport N76543210: interview appointment on 16 October 2026.','message_id':'test-appt','attachments':[]}
    with patch('app.email_agent.smtp_imap.fetch_unseen',return_value=[apptmsg]):
        with app.app_context(): result=process_inbox()
    with app.app_context():
        drafts=Applicant.query.filter_by(passport_no='N76543210',kind='EXTEND').all(); assert drafts
        draft=drafts[-1]; assert (draft.departure_date,draft.arrival_date)==(date(2026,10,30),date(2027,4,30))
    assert result['appointments']==1
    print('PASS simulated inbox policy mapping + embassy appointment → EXTEND draft (+14d/+6mo)')

    # MCP tools share the same date/city/police pipeline.
    from mcp_server.server import calculate_dates, match_city, extract_police_report
    md=calculate_dates('NEW','31/08/2026',company='ALLIANZ')
    assert md['arrival_date']=='2027-02-28' and md['arrival_date_formatted']=='02/28/2027',md
    assert match_city('SAMAGI MAWATHA KATUKENDA WEST, DANKOTUWA, SRI LANKA')['city']=='Dankotuwa'
    mg=extract_police_report(str(POLICE1)); assert mg['city']=='Dankotuwa' and mg['addr_no']=='41/A',mg
    print('PASS MCP date/city tools + universal police extraction')
    print('SYSTEM CHECK RESULT: all isolated functional checks passed; no real email was sent and no live database records were changed.')

if __name__=='__main__': main()
